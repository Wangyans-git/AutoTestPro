#!C:\wys\AutoTestProjects
# -*- coding: utf-8 -*-
# @Time    :
# @Author  :yansheng.wang
# @File    :
# @Description : OTA次数统计

import os
import re
import threading
import time
from datetime import datetime

import openpyxl
import serial

from H7135.logs.get_log import GetLog

log_name = datetime.now().strftime("logs\\iot_test_%Y%m%d_%H%M%S.txt")
time_log = datetime.now().strftime("_%Y%m%d_%H%M%S.txt")


class H7135_OTA(object):

    def __init__(self, com, dbs, com1, dbs1, timeout, update_count):
        self.com = com
        self.dbs = dbs
        self.timeout = timeout
        self.com1 = com1  # 继电器
        self.dbs1 = dbs1  # 继电器
        self.success_count = 0  # 成功OTA次数
        self.start_count = 0  # 开始OTA次数
        self.sensor_row = 0  # 记录sensor行数
        self.sensor_send_count = 0  # 记录sensor发送次数
        self.update_count = update_count  # OTA次数
        self.start_time = None
        self.end_time = None
        self.get_log_send = GetLog(f"logs\\发送_log.txt")
        self.get_log_recv = GetLog(f"logs\\接收_log.txt")
        self.filename = openpyxl.Workbook()
        self.sheet = self.filename.create_sheet(index=0, title="OTA成功率")
        self.sheet1 = self.filename.create_sheet(index=0, title="分布式网关成功率")
        self.folder_path = os.path.dirname(os.path.abspath(__file__))
        self.result_file = r'{}\report\{}_{}.xlsx'.format(self.folder_path, "H7135压测结果",
                                                          datetime.now().strftime('%Y-%m-%d_%H.%M.%S.%f'))
        try:
            self.ser = serial.Serial(self.com,
                                     self.dbs,
                                     timeout=self.timeout)
            print("*********打开日志串口成功*********")
        except Exception as e:
            print("*********日志串口异常:{}*********".format(e))
        try:
            self.ser_Relay = serial.Serial(self.com1,
                                           self.dbs1,
                                           timeout=self.timeout)
            print("*********打开继电器串口成功*********")
        except Exception as e:
            print("*********继电器串口异常:{}*********".format(e))

    # 读取处理数据
    def read_date(self):
        self.sheet.cell(1, 1).value = "序号"
        self.sheet.cell(1, 2).value = "开始时间"
        self.sheet.cell(1, 3).value = "结束时间"
        self.sheet.cell(1, 4).value = "是否成功"
        self.sheet.cell(1, 5).value = "用时"
        self.sheet.cell(1, 6).value = "当前总成功率"
        self.sheet.cell(1, 7).value = "Ota_DoCalculate_MD5"

        self.sheet1.cell(1, 1).value = "序号"
        self.sheet1.cell(1, 2).value = "发送广播"
        self.sheet1.cell(1, 3).value = "接收广播"
        self.sheet1.cell(1, 4).value = "当前总成功率"
        row_num = 0
        sensor_recv_count = 0
        self.thread_sensor()  # 发送sensor广播
        with open(f'logs\\串口日志记录{time_log}', 'w') as file_handle:
            while row_num <= self.update_count:
                try:
                    date_line = self.ser.readline().decode()
                    file_handle.write(date_line)
                    match = re.search(r'\[[a-fA-F0-9]{32}]', date_line)
                    """
                    分布式网关处理
                    """
                    # if "RT/gateway/group/data/report/GD/" in date_line:  # 新固件断言
                    if "insert trigger sensor data is OK" in date_line:
                        sensor_recv_count += 1
                        print(f"收到sensor信息次数:{sensor_recv_count}次")
                        self.get_log_recv.info(f"收到sensor信息次数:{sensor_recv_count}次")
                        self.sheet1.cell(self.sensor_row + 1, 3).value = "接收成功"
                        self.sheet1.cell(self.sensor_row + 1,
                                         4).value = f"{sensor_recv_count / self.sensor_send_count * 100:.2f}%"
                        self.filename.save(self.result_file)
                    """
                    OTA处理
                    """
                    if "Number of blocks remaining:" in date_line:
                        print(date_line)

                    if "Number of blocks remaining: 891" in date_line:
                        row_num += 1
                        print("开始")
                        self.start_time = datetime.now()
                        self.start_count += 1
                        self.sheet.cell(row_num + 1, 1).value = f"第{row_num}次"
                        self.sheet.cell(row_num + 1, 2).value = self.start_time
                        self.filename.save(self.result_file)
                    elif match:
                        print(f"Ota_DoCalculate_MD5:{match.group(0)} ")
                        self.sheet.cell(row_num + 1, 7).value = match.group(0)
                        if "[98702e10db4c9d049548549d8237e072]" in date_line:
                            print("结束")
                            self.success_count += 1
                            self.end_time = datetime.now()
                            start_to_end_time = (self.end_time - self.start_time)
                            self.sheet.cell(row_num + 1, 3).value = self.end_time
                            self.sheet.cell(row_num + 1, 4).value = "成功"
                            self.sheet.cell(row_num + 1, 5).value = start_to_end_time
                            print(f"OTA用时：{start_to_end_time}")
                            print(f"{self.success_count / self.start_count * 100:.2f}%")
                            self.sheet.cell(row_num + 1,
                                            6).value = f"{self.success_count / self.start_count * 100:.2f}%"
                            self.filename.save(self.result_file)
                        else:
                            self.sheet.cell(row_num + 1, 4).value = "失败"
                            self.sheet.cell(row_num + 1,
                                            6).value = f"{self.success_count / self.start_count * 100:.2f}%"
                            self.filename.save(self.result_file)
                except Exception as e:
                    # print(e)
                    pass
            self.ser.close()
        print(
            f"一共OTA{self.update_count}次，成功{self.success_count}次,成功率{self.success_count / self.update_count * 100:.2f}%")

    def sensor_check(self):
        time.sleep(5)
        while True:
            try:
                self.ser_Relay.write(bytes.fromhex('A0 01 01 A2'))
                time.sleep(0.5)
                self.ser_Relay.write(bytes.fromhex('A0 01 00 A1'))
                self.sensor_send_count += 1
                self.sensor_row += 1
                self.sheet1.cell(self.sensor_row + 1, 1).value = f"第{self.sensor_row}次"
                self.sheet1.cell(self.sensor_row + 1, 2).value = "发送成功"
                self.filename.save(self.result_file)
            except Exception as e:
                print(e)
                pass
            print(f"发送sensor广播次数:{self.sensor_send_count}次")
            self.get_log_send.info(f"发送sensor广播次数:{self.sensor_send_count}次")
            time.sleep(10)

    def thread_sensor(self):
        thread_ = threading.Thread(target=self.sensor_check)
        thread_.start()


if __name__ == '__main__':
    start_ota_count = 1000
    program = H7135_OTA('com3', 115200, 'com4', 9600, 3, start_ota_count)
    program.read_date()
