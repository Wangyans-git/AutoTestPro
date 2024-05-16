import os
import re
import sys
from datetime import datetime
from time import gmtime
from time import strftime

import openpyxl as openpyxl
from loguru import logger

folder_path = ""
if getattr(sys, 'frozen', False):
    folder_path = os.path.dirname(sys.executable)
elif __file__:
    folder_path = os.path.dirname(__file__)

print(folder_path)


class Export_Raport:
    # 创建报告存放文件夹
    if not os.path.exists('report/'):
        os.makedirs('report/')
    logfilename = []
    interrupt_time = 0
    connection_router_time = 0

    # 获取.log的文件路径和文件名
    def get_file(self, folder_path):

        dir_file = os.listdir(folder_path)

        for path in dir_file:
            if path[-3:] == 'log':
                # whole_path = r'{}//{}'.format(folder_path, path)
                whole_pathname = path[:-4]
                # logger.info('chuangjian: {}'.format(path))
                # self.logfile.append(whole_path)
                self.logfilename.append(whole_pathname)
        return self.logfilename

    def get_interrupt(self):
        for i in self.logfilename:
            whole_path = r'{}\{}.log'.format(folder_path, i)

            f = open(whole_path, 'r+', encoding='utf-8', errors='ignore')
            line = f.readline()

            while line:
                if '*****read system_data********' in line:
                    interrupt = line
                    self.interrupt_time = self.interrupt_time + 1

                if 'start to connect ssid:' in line:
                    connection_router = line
                    self.connection_router_time = self.connection_router_time + 1

                line = f.readline()
            f.close()
            logger.info('interrupt_time: {}'.format(self.interrupt_time))
        return self.interrupt_time, self.connection_router_time

    def Log_extraction(self, reportname, breakkey, startkey, endkey):
        self.breakkey = breakkey
        self.startkey = startkey
        self.endkey = endkey

        j = 0
        filecount = 0
        self.startconnectwifikey = 'start to connect ssid:'
        self.startconnectiotkey = 'HAL_TLS_Connect start...'

        self.filename = openpyxl.Workbook()
        sheett = self.filename.create_sheet(index=1, title="report")
        result_file = r'{}\report\{}_{}.xlsx'.format(folder_path, reportname,
                                                     datetime.now().strftime('%Y-%m-%d_%H.%M.%S.%f'))
        for i in self.logfilename:
            listtime = []
            whole_path = f"{folder_path}\{i}.log"
            j = j + 1

            # # [2022-11-18_19:54:13:832]
            self.yre_str = r'[0-2][0-9][0-3][0-9]-[0-2][0-9]-[0-3][0-9]_[0-5][0-9]:[0-5][0-9]:[0-5][0-9]:[0-9]{3}'
            self.ytime_format = '%Y-%m-%d_%H:%M:%S:%f'

            # [[20230707_19:11:52:669]]
            self.jre_str = r'[0-2][0-9][0-3][0-9][0-2][0-9][0-3][0-9]_[0-5][0-9]:[0-5][0-9]:[0-5][0-9]:[0-9]{3}'
            self.jtime_format = '%Y%m%d_%H:%M:%S:%f'
            # self.Raport_BASE_NAME = "LOG"

            # [20230116_12:16:20:826]
            self.cre_str = r'[0-2][0-9][0-3][0-9][0-2][0-9][0-3][0-9]_[0-5][0-9]:[0-5][0-9]:[0-5][0-9]:[0-9]{3}'
            self.ctime_format = '%Y%m%d_%H:%M:%S:%f'

            sheet = self.filename.create_sheet(index=1, title=i)
            # sheet = '{}_{}'.format(sheet, j)
            sheet.cell(1, 1).value = '序号'
            sheet.cell(1, 2).value = 'begin'
            sheet.cell(1, 3).value = 'end'
            sheet.cell(1, 4).value = '时间差(分秒)'
            sheet.cell(1, 5).value = '时间差(秒)'

            # print(result_file)
            logger.info('result_file: {}'.format(result_file))
            stack_overflow_time = 0
            Crash_time = 0
            crash_time = 0
            rtw_recv_time = 0
            Startcount = 0
            Successcount = 0
            break_times = 0
            startconnectwifi_times = 0
            startconnectiot_times = 0
            endconnectwifi_times = 0

            f = open(whole_path, 'r+', encoding='utf-8', errors='ignore')
            line = f.readline()
            line1 = None
            line2 = None
            line3 = None
            linesku = None
            lineuuid = None
            linewifimac = None
            linewifi_hard_ver = None
            linewifi_soft_ver = None
            while line:

                if 'stack overflow' in line:
                    stack_overflow = line
                    stack_overflow_time += 1

                if 'Crash' in line:
                    Crash = line
                    Crash_time += 1
                if 'crash' in line:
                    crash = line
                    crash_time += 1
                if 'rtw_recv_' in line:
                    rtw_recv = line
                    rtw_recv_time += 1

                if 'sku:' in line:
                    linesku = line
                if 'uuid:' in line:
                    lineuuid = line

                if 'wifi_mac:' in line:
                    linewifimac = line
                if 'wifi_hard_ver:' in line:
                    linewifi_hard_ver = line
                if 'wifi_soft_ver:' in line:
                    linewifi_soft_ver = line

                if self.breakkey in line:
                    linebreak = line
                    break_times = break_times + 1
                    logger.info('break_times:[{}] {}'.format(break_times, linebreak))
                if self.startconnectwifikey in line:
                    startconnectwifi = line
                    startconnectwifi_times = startconnectwifi_times + 1
                    logger.info('startconnectwifi_times:[{}] {}'.format(startconnectwifi_times, startconnectwifi))

                if 'WIFI  wlan0 Setting:' in line:
                    endconnectwifi_times = endconnectwifi_times + 1

                if self.startconnectiotkey in line:
                    startconnectiot = line
                    startconnectiot_times = startconnectiot_times + 1
                    logger.info('startconnectiot_times:[{}] {}'.format(startconnectiot_times, startconnectiot))

                if self.startkey in line:
                    line1 = line
                    Startcount += 1
                    sheet.cell(Startcount + 1, 1).value = Startcount
                    sheet.cell(Startcount + 1, 2).value = line1
                    # print(line)

                if self.startkey in line:

                    line1 = line
                    # Startcount += 1
                    # print(line)
                elif line1 and self.endkey in line:
                    line2 = line
                    # print(line2)
                    logger.info('[{}] {}'.format(Successcount, line1))
                    logger.info('[{}] {}'.format(Successcount, line2))
                    Successcount += 1
                    # sheet.cell(Startcount + 1, 1).value = Startcount
                    # sheet.cell(Startcount + 1, 2).value = line1
                    sheet.cell(Startcount + 1, 3).value = line2
                    time1_str = re.search(self.yre_str, line1).group()
                    time1 = datetime.strptime(time1_str, self.ytime_format)
                    time2_str = re.search(self.yre_str, line2).group()
                    time2 = datetime.strptime(time2_str, self.ytime_format)
                    delta = time2 - time1
                    time_str = delta.total_seconds()
                    listtime.append(time_str)

                    time_M_S = strftime("%M:%S", gmtime(time_str))
                    sheet.cell(Startcount + 1, 4).value = time_M_S
                    sheet.cell(Startcount + 1, 5).value = time_str
                    line1 = None
                    line3 = None
                line = f.readline()
            f.close()

            if break_times == 0:
                logger.info('没开始')
            else:
                if not listtime:
                    logger.info('listt:{}'.format(listtime))
                else:
                    # 最小值
                    listtime.sort()
                    minimum = listtime[0]
                    # if reportname == '上电-wifi':
                    #     Greater_than = 20
                    # elif reportname == '上电-iot':
                    #     Greater_than = 20
                    # elif reportname == 'wifi-iot':
                    #     Greater_than = 30
                    Greater_than = 20
                    lessmiute = 0
                    for minute in listtime:
                        if minute > Greater_than:
                            lessmiute = lessmiute + 1
                        else:
                            pass
                    # 最大值
                    listtime.sort(reverse=True)
                    Maximum = listtime[0]
                    sum = 0
                    # 平均值，这一块有点小问题
                    for item in listtime:
                        sum += item
                        average = sum / len(listtime)
                    # 成功率
                    Success_rate = Successcount / break_times
                    Success_rate = '{:.2%}'.format(Success_rate)
                    if len(listtime) == break_times:
                        pass
                    else:
                        # Maximum = 200
                        falsetime = break_times - len(listtime)
                        average1 = (sum + falsetime * 200) / break_times

                    connectiotfailurestimes = startconnectiot_times - Successcount
                    connectwififailurestimes = startconnectwifi_times - Successcount

                    if connectiotfailurestimes < 0:
                        connectiotfailurestimes = 0
                    else:
                        pass
                    if connectwififailurestimes < 0:
                        connectwififailurestimes = 0
                    else:
                        pass

                    if linesku != None and lineuuid != None and linewifimac != None and linewifi_hard_ver != None and linewifi_soft_ver != None:
                        linesku = linesku[-6:]
                        lineuuid = lineuuid[-24:]
                        linewifimac = linewifimac[-18:]
                        linewifi_hard_ver = linewifi_hard_ver[-8:]
                        linewifi_soft_ver = linewifi_soft_ver[-8:]
                    else:
                        pass
                    # 表头信息
                    Greater_other = '大于{}秒个数'.format(Greater_than)
                    sheett.cell(1, 1).value = '序号'
                    sheett.cell(1, 2).value = reportname
                    sheett.cell(1, 3).value = '开始关键字'
                    sheett.cell(1, 4).value = '结束关键字'
                    sheett.cell(1, 5).value = '成功次数'
                    sheett.cell(1, 6).value = '最小值(秒)'
                    sheett.cell(1, 7).value = '最大值(秒)'
                    sheett.cell(1, 8).value = Greater_other
                    sheett.cell(1, 9).value = '平均值(秒)'

                    sheett.cell(1, 11).value = '重启关键字'
                    sheett.cell(1, 12).value = '结束关键字'
                    sheett.cell(1, 13).value = '开始次数'
                    sheett.cell(1, 14).value = '成功次数'
                    sheett.cell(1, 15).value = '成功率'

                    #
                    if reportname in ['设备断电上电-wifi', '设备断电上电-iot']:
                        sheett.cell(1, 17).value = '连接wifi次数'
                        sheett.cell(1, 18).value = '连接wifi失败次数'
                        sheett.cell(1, 19).value = '连接iot次数'
                        sheett.cell(1, 20).value = '连接iot失败次数'
                    else:
                        pass
                    if linesku != None and lineuuid != None and linewifimac != None and linewifi_hard_ver != None and linewifi_soft_ver != None:
                        sheett.cell(1, 23).value = 'sku'
                        sheett.cell(1, 24).value = 'mac'
                        sheett.cell(1, 25).value = 'hard_version'
                        sheett.cell(1, 26).value = 'soft_version'
                        sheett.cell(1, 27).value = 'wifi_mac'
                    else:
                        pass

                    sheett.cell(1, 28).value = 'stack_overflow'
                    sheett.cell(1, 29).value = 'Crash'
                    sheett.cell(1, 30).value = 'rtw_recv'

                    filecount += 1
                    sheett.cell(filecount + 1, 1).value = filecount  # 序号
                    sheett.cell(filecount + 1, 2).value = i
                    sheett.cell(filecount + 1, 3).value = self.startkey
                    sheett.cell(filecount + 1, 4).value = self.endkey
                    sheett.cell(filecount + 1, 5).value = len(listtime)
                    sheett.cell(filecount + 1, 6).value = minimum
                    sheett.cell(filecount + 1, 7).value = Maximum
                    sheett.cell(filecount + 1, 8).value = lessmiute + break_times - Successcount
                    sheett.cell(filecount + 1, 9).value = average

                    sheett.cell(filecount + 1, 11).value = self.breakkey
                    sheett.cell(filecount + 1, 12).value = self.endkey
                    sheett.cell(filecount + 1, 13).value = break_times
                    sheett.cell(filecount + 1, 14).value = Successcount
                    sheett.cell(filecount + 1, 15).value = Success_rate

                    if reportname in ['设备断电上电-wifi', '设备断电上电-iot']:
                        sheett.cell(filecount + 1, 17).value = startconnectwifi_times
                        sheett.cell(filecount + 1, 18).value = startconnectwifi_times - endconnectwifi_times
                        sheett.cell(filecount + 1, 19).value = startconnectiot_times
                        sheett.cell(filecount + 1, 20).value = connectiotfailurestimes
                    else:
                        pass

                    if linesku != None and lineuuid != None and linewifimac != None and linewifi_hard_ver != None and linewifi_soft_ver != None:
                        sheett.cell(filecount + 1, 23).value = linesku
                        sheett.cell(filecount + 1, 24).value = lineuuid
                        sheett.cell(filecount + 1, 25).value = linewifi_hard_ver
                        sheett.cell(filecount + 1, 26).value = linewifi_soft_ver
                        sheett.cell(filecount + 1, 27).value = linewifimac
                    else:
                        pass
                    sheett.cell(filecount + 1, 28).value = stack_overflow_time
                    sheett.cell(filecount + 1, 29).value = Crash_time + crash_time
                    sheett.cell(filecount + 1, 30).value = rtw_recv_time

                    self.filename.save(result_file)
        return result_file

    def distinction(self):
        interrupt_time = Export_Raport().get_interrupt()[0]
        connection_router_time = Export_Raport().get_interrupt()[1]
        if interrupt_time > 5:
            Export_Raport().Log_extraction('设备断电上电-iot', '*****read system_data********',
                                           '*****read system_data********',
                                           'mqtt_service_subscribe_topic subscribe success!')
            Export_Raport().Log_extraction('设备断电上电-wifi', '*****read system_data********',
                                           '*****read system_data********', 'wifi station: got ip')
        elif connection_router_time > 10:
            Export_Raport().Log_extraction('路由器上电-iot', ']MQTT disconnect.', 'wifi:connected with',
                                           'wifi station: got ip')
        else:
            Export_Raport().Log_extraction('断网上电-iot ', 'event_handler(100):MQTT disconnect.',
                                           'HAL_TLS_Connect start...',
                                           'mqtt_service_subscribe_topic subscribe success!')
        logger.info('interrupt_timez: {}'.format(interrupt_time))


if __name__ == '__main__':
    Export_Raport().get_file(folder_path=folder_path)
    # '*****read system_data********' , 'Cloud Device Construct Success!','mqtt_service_subscribe_topic subscribe success!'
    # Export_Raport().Log_extraction('上电-wifi','*****read device_data********','HAL_Wifi_Connect(253):start to connect ssid:','Cloud Device Construct Success!')
    # Export_Raport().distinction()
    # Export_Raport().Log_extraction('路由器上电-iot', 'ap_probe_send over, resett wifi status to disassoc', 'start to connect ssid', 'Cloud Device Construct Success!')
    # _disconnect_wifi(155):ret,mqtt disconnect!
    Export_Raport().Log_extraction('设备断电上电-iot', 'Rtl8735b IoT Platform', 'Rtl8735b IoT Platform',
                                   'subscribe success!')
    # Export_Raport().Log_extraction('路由器断网-iot ', 'mqtt disconnect!', 'start to connect ssid',
    #                                'Cloud Device Construct Success!')

    # Export_Raport().Log_extraction('上电-wifi','MQTT Disconnect','MQTT Disconnect','IP              =>')
    # Export_Raport().Log_extraction('上电-iot','Boot Loader <==','Boot Loader <==','Receive Message With topicName:[GD/')
    # Export_Raport().Log_extraction('上电-wifi','Boot Loader <==','Boot Loader <==','WiFi station interface connected')
    # Export_Raport().Log_extraction('上电-iot','Boot Loader <==','Boot Loader <==','Received command for Endpoint=1 Cluster=0x0000_0006 Command=0x0000_000')
    # Export_Raport().Log_extraction('上电-iot','read system_data','read system_data','Ihoment:MQTT Connected')
    # Export_Raport().Log_extraction('上电-wifi','read system_data','read system_data','set group key to hw: alg:4(WEP40-1 WEP104-5 TKIP-2 AES-4) keyid:1')
    # Export_Raport().Log_extraction( 'wifi-iot','read system_data','set group key to hw: alg:4(WEP40-1 WEP104-5 TKIP-2 AES-4) keyid:1','Ihoment:MQTT Connected')
    # Export_Raport().Log_extraction('上电-iot','read system_data','read system_data','mqtt_service_subscribe_topic subscribe success!')
    # 'event_handler(100):MQTT disconnect.',
    # Export_Raport().Log_extraction('路由器上电-iot', '[Driver]: no beacon for a long time, disconnect or roaming',
    #                                'start to connect ssid:',
    #                                'Cloud Device Construct Success!')
    # Export_Raport().Log_extraction('wifi-iot','read system_data','WIFI  wlan0 Setting:','mqtt_service_subscribe_topic subscribe success!')
    # Export_Raport().Log_extraction('上电-iot','Model: AXERA AX620_demo Board','Model: AXERA AX620_demo Board','mqtt_service_subscribe_topic subscribe success!')
    # Export_Raport().Log_extraction('上电-wifi','Model: AXERA AX620_demo Board','Model: AXERA AX620_demo Board','Connection to b0:be:76:0b:2a:07 completed')
    # Export_Raport().Log_extraction('wifi-iot','Model: AXERA AX620_demo Board','Connection to b0:be:76:0b:2a:07 completed','mqtt_service_subscribe_topic subscribe success!')

    # Export_Raport().Log_extraction('路由器上电-iot', 'disconnect: 0X', 'wifi station: got ip')

    # Export_Raport().Log_extraction('wifi连接-iot', 'ble_config_event_handler(290): IOT_BLE_WIFI_CONFIG_START',
    #                                'start to connect ssid:',
    #                                'Cloud Device Construct Success!')

# print(whole_path)
# breakkey =
# startkey = # disconnect MQTT for some reasons..
# "mqtt will destroy"

# breakkey = "Boot Loader <=="


# startkey= ":needupdate: true,"
# "Boot Loader <=="
# HAL_TLS_Connect start...
# endkey = "set group key to hw: alg:4(WEP40-1 WEP104-5 TKIP-2 AES-4) keyid:1"
# mqtt_service_subscribe_topic subscribe success!
# endkey = "device will reboot because of ota"
